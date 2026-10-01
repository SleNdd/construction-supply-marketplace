"""Собирает пояснительную записку из content.md."""

from pathlib import Path
import argparse
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Пояснительная_записка_ОбъектМаркет.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    tc_pr.append(borders)


def set_cell_margins(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "bottom", "left", "right"):
        tag = OxmlElement(f"w:{side}")
        tag.set(qn("w:w"), "90" if side in ("top", "bottom") else "110")
        tag.set(qn("w:type"), "dxa")
        margins.append(tag)
    tc_pr.append(margins)


def add_field(paragraph, instruction, fallback=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for part in (begin, code, separate, result, end):
        run._r.append(part)


def add_inline(paragraph, text):
    pieces = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(11)
        elif piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        else:
            paragraph.add_run(piece)


TABLE_TITLES = {
    1: "Сравнение действующих решений",
    2: "Права участников системы",
    3: "Основные сущности базы данных",
    4: "Контрольные примеры расчётов",
    5: "Основные группы методов API",
    6: "Сравнение вариантов закупки цемента в учебном наборе",
    7: "Параметры конфигурации стенда",
    8: "Показатели для сравнительного испытания",
    9: "Сценарии проверки системы",
}


def add_table(document, rows, index):
    headers = [item.strip() for item in rows[0].strip("|").split("|")]
    body = rows[1:]
    if body and re.fullmatch(r"\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*", body[0]):
        body = body[1:]
    caption = document.add_paragraph(style="Caption")
    caption.add_run(f"Таблица {index} — ")
    caption.add_run(TABLE_TITLES.get(index, "Данные испытаний"))
    caption.paragraph_format.keep_with_next = True
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    compact = index == 9
    for row_index, data in enumerate([headers] + [[item.strip() for item in row.strip("|").split("|")] for row in body]):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for cell, value in zip(cells, data):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0 if compact else 1.15
            add_inline(p, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "D9E4EE")
                for run in p.runs:
                    run.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F6F8FA")
            for run in p.runs:
                if run.font.name != "Consolas":
                    run.font.name = "Times New Roman"
                run.font.size = Pt(10 if compact else 11)
    for cell in table.rows[0].cells:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)
        break
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def create(output=OUTPUT):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)
    section.different_first_page_header_footer = True

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.widow_control = True

    for name, size, before, after in [("Heading 1", 14, 14, 8), ("Heading 2", 14, 12, 5)]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.line_spacing = 1.15
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = Pt(18)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    styles["Caption"].font.name = "Times New Roman"
    styles["Caption"].font.size = Pt(12)
    styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Caption"].paragraph_format.first_line_indent = Cm(0)
    styles["Caption"].paragraph_format.space_before = Pt(8)
    styles["Caption"].paragraph_format.space_after = Pt(4)
    styles["Caption"].paragraph_format.line_spacing = 1.0
    for name, indent in (("TOC 1", 0), ("TOC 2", 0.55)):
        toc = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        toc.font.name = "Times New Roman"
        toc.font.size = Pt(10)
        toc.font.color.rgb = RGBColor(0, 0, 0)
        toc.paragraph_format.line_spacing = 1.0
        toc.paragraph_format.space_before = Pt(0)
        toc.paragraph_format.space_after = Pt(0)
        toc.paragraph_format.first_line_indent = Cm(0)
        toc.paragraph_format.left_indent = Cm(indent)
        toc.paragraph_format.keep_with_next = False
        toc.paragraph_format.keep_together = False
        toc.paragraph_format.widow_control = False

    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.first_line_indent = Cm(0)
    cover.paragraph_format.space_after = Pt(0)
    cover.add_run("Астраханский государственный\nархитектурно-строительный университет\n").bold = True
    cover.add_run("Направление подготовки 09.03.02\nИнформационные системы и технологии")
    document.add_paragraph("\n\n\n")
    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    title_run = p.add_run("ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА")
    title_run.font.name = "Times New Roman"
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Разработка информационной системы выбора и поставки строительных материалов с учётом потребностей строительного объекта")
    document.add_paragraph("\n\n")
    p = document.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Выполнил: студент группы ЗИТ-51-22\nБальдюсов Кирилл Андреевич\n\n")
    p.add_run("Руководитель: ______________________________\n")
    p.add_run("Кафедра: ___________________________________")
    document.add_paragraph("\n\n\n")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Астрахань 202_")
    document.add_page_break()

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, " PAGE ", "1")

    lines = (ROOT / "content.md").read_text(encoding="utf-8").splitlines()
    table_index = 0
    figure_index = 0
    i = 0
    in_bibliography = False
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        figure = re.fullmatch(r"!\[([^\]]+)\]\(([^)]+)\)", line)
        if figure:
            caption_text, relative_path = figure.groups()
            image_path = (ROOT / relative_path).resolve()
            if not image_path.is_file():
                raise FileNotFoundError(image_path)
            figure_index += 1
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
            width = Cm(8) if "mobile" in image_path.stem else Cm(16.5)
            shape = paragraph.add_run().add_picture(str(image_path), width=width)
            shape._inline.docPr.set("descr", caption_text)
            caption = document.add_paragraph(style="Caption")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.first_line_indent = Cm(0)
            caption.paragraph_format.keep_with_next = False
            caption.add_run(f"Рисунок {figure_index} — {caption_text}")
            continue
        if line.startswith("|"):
            rows = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            table_index += 1
            add_table(document, rows, table_index)
            continue
        if line.startswith("# "):
            title = line[2:]
            if title == "СОДЕРЖАНИЕ":
                heading = document.add_paragraph()
                heading.paragraph_format.page_break_before = True
                heading.paragraph_format.first_line_indent = Cm(0)
                heading.paragraph_format.space_after = Pt(8)
                heading.add_run(title).bold = True
                p = document.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0)
                add_field(p, ' TOC \\o "1-2" \\h \\z \\u ', "Содержание обновляется в Word")
                while i < len(lines) and not lines[i].strip():
                    i += 1
                if i < len(lines) and lines[i].startswith("Содержание формируется"):
                    i += 1
                continue
            if title == "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ":
                in_bibliography = True
            heading = document.add_heading(title, level=1)
            if title in {"ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"} or title.startswith(("1 ", "2 ", "3 ", "4 ", "ПРИЛОЖЕНИЕ")):
                heading.paragraph_format.page_break_before = True
            continue
        if line.startswith("## "):
            document.add_heading(line[3:], level=2)
            continue
        p = document.add_paragraph()
        if in_bibliography and re.match(r"^\d+\. ", line):
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.space_after = Pt(0)
        add_inline(p, line)

    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    document.core_properties.title = "Разработка информационной системы выбора и поставки строительных материалов с учётом потребностей строительного объекта"
    document.core_properties.author = "Бальдюсов Кирилл Андреевич"
    document.core_properties.last_modified_by = "Бальдюсов Кирилл Андреевич"
    document.core_properties.comments = ""
    document.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Сборка пояснительной записки")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    create(parser.parse_args().output)
